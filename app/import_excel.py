"""Liest `Produktionsplanung_V7_3.xlsm` in die Datenbank.

Gelesen werden nur die Eingabespalten. Alles, was das Excel rechnet (Bezeichnung,
Material, Maschine, Zeiten, Liefer-KW, Plan-KW, KW eff., Rohmaterialmengen), rechnet die
App aus den Regeln neu; so fallen die uneinheitlichen Stellen aus `docs/regeln.md` von
selbst weg.

Der Import prüft und meldet, er korrigiert nicht still (Hausregel):

- Materialbezeichnungen, die nicht in der Liste stehen (R-010)
- Auftragszeilen mit unbekannter Artikelnummer (R-011)
- Statuswerte, die nicht `offen`, `in Arbeit` oder `erledigt` heissen (R-008)
- Rohmaterialien über die 15 Zeilen hinaus, die das Blatt `Bedarf` kennt
- Kundennamen werden anonymisiert ("Kunde 01", "Kunde 02", ...), echte Namen landen nie
  in der Datenbank dieser App.

Aufruf:

    uv run python -m app.import_excel tests/daten/Produktionsplanung_demo.xlsm
    uv run python -m app.import_excel <datei> --trocken   (nur prüfen, nichts schreiben)
"""

from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.datenbank import Sitzung, tabellen_anlegen
from app.modelle import (
    Artikel,
    Auftrag,
    Lieferant,
    Maschine,
    Parametersatz,
    Rohmaterial,
    Stuecklistenposition,
)
from app.regeln.parameter import PARAMETER_KATALOG
from app.regeln.stammdaten import material_ist_bekannt
from app.regeln.typen import Kalenderwoche, Status

# Bezeichnung im Blatt `Parameter` -> Name in der Tabelle `parameter` (R-001).
PARAMETERNAMEN = {
    "Sicherheitsbestand %": "sicherheitsbestand_prozent",
    "Auslastungsgrenze": "auslastungsgrenze_prozent",
    "Rüstzuschlag Alu (min)": "ruestzuschlag_alu_min",
    "Losgrenze Alu (Stk)": "losgrenze_alu_stk",
    "Schichten pro Tag": "schichten_pro_tag",
    "Std pro Schicht": "stunden_pro_schicht",
    "Arbeitstage pro Woche": "arbeitstage_pro_woche",
    "Puffer Lieferant C (Wochen)": "puffer_lieferant_c_wochen",
    "Kostenstelle Einkauf": "kostenstelle_einkauf",
    "Planungshorizont (Wochen)": "planungshorizont_wochen",
    "Wartungsabzug (h)": "wartungsabzug_h",
    "Vorlauf Prio A (Wochen)": "vorlauf_prio_a_wochen",
}

# Diese beiden Parameter gibt es im Excel nicht: dort steht die 1 fest in der Formel
# für Plan-KW (R-006, offene Frage F-3).
ERGAENZTE_PARAMETER = {
    "vorlauf_prio_b_wochen": Decimal(1),
    "vorlauf_prio_c_wochen": Decimal(1),
}

# Diese Parameter übernimmt die App nicht: Export-Pfad ist Konfiguration (F-6),
# aktuelle KW und Montag kommen aus dem Datum (R-002), der Edelstahl-Zuschlag ist
# seit V5 tot (Abschnitt 9 des Katalogs).
NICHT_UEBERNOMMEN = {
    "Aktuelle KW",
    "Jahr",
    "Montag aktuelle KW",
    "Export-Pfad",
    "Zuschlag Edelstahl (min)",
}


@dataclass
class Importbericht:
    """Was gelesen wurde und was aufgefallen ist."""

    gelesen: dict[str, int] = field(default_factory=dict)
    meldungen: list[str] = field(default_factory=list)

    def zaehle(self, was: str, anzahl: int) -> None:
        self.gelesen[was] = anzahl

    def melde(self, text: str) -> None:
        self.meldungen.append(text)

    def als_text(self) -> str:
        zeilen = [f"{was}: {anzahl}" for was, anzahl in self.gelesen.items()]
        if self.meldungen:
            zeilen.append("")
            zeilen.append(f"{len(self.meldungen)} Meldungen:")
            zeilen += [f"  - {text}" for text in self.meldungen]
        else:
            zeilen.append("")
            zeilen.append("Keine Meldungen.")
        return "\n".join(zeilen)


def _text(wert: object) -> str:
    return "" if wert is None else str(wert).strip()


def _zahl(wert: object) -> Decimal:
    if wert is None or _text(wert) == "":
        return Decimal(0)
    return Decimal(str(wert))


def _prozent(wert: object) -> Decimal:
    """Excel führt Prozentzellen als Anteil (0,85). Die App speichert 85."""
    zahl = _zahl(wert)
    return zahl * 100 if zahl <= 1 else zahl


def _datum(wert: object) -> date | None:
    if isinstance(wert, datetime):
        return wert.date()
    if isinstance(wert, date):
        return wert
    return None


def lies_excel(datei: Path, sitzung: Session) -> Importbericht:
    """Liest die Arbeitsmappe und schreibt die Stammdaten und Aufträge."""
    from openpyxl import load_workbook

    bericht = Importbericht()
    mappe = load_workbook(datei, data_only=True, read_only=True)

    _parameter(mappe["Parameter"], sitzung, bericht)
    jahr = _planungsjahr(mappe["Parameter"])
    _maschinen(mappe["Kapazität"], sitzung, bericht, jahr)
    _lieferanten(mappe["Lieferanten"], sitzung, bericht)
    _rohmaterial(mappe["Rohmaterial"], sitzung, bericht)
    artikelnummern = _artikel(mappe["Artikel"], sitzung, bericht)
    sitzung.flush()
    _stueckliste(mappe["Stückliste"], sitzung, bericht, artikelnummern)
    _auftraege(mappe["Aufträge"], sitzung, bericht, artikelnummern, jahr)

    mappe.close()
    return bericht


def _planungsjahr(blatt) -> int:
    for zeile in blatt.iter_rows(min_row=1, max_row=30, max_col=2, values_only=True):
        if _text(zeile[0]) == "Jahr":
            return int(_zahl(zeile[1]))
    return date.today().year


def _parameter(blatt, sitzung: Session, bericht: Importbericht) -> None:
    gelesen = 0
    for zeile in blatt.iter_rows(min_row=2, max_row=30, max_col=2, values_only=True):
        bezeichnung = _text(zeile[0])
        if not bezeichnung or bezeichnung in NICHT_UEBERNOMMEN:
            continue
        name = PARAMETERNAMEN.get(bezeichnung)
        if name is None:
            bericht.melde(f"Parameter '{bezeichnung}' ist unbekannt, nicht übernommen")
            continue
        wert = _prozent(zeile[1]) if name.endswith("_prozent") else _zahl(zeile[1])
        beschriftung, einheit = PARAMETER_KATALOG[name]
        sitzung.merge(
            Parametersatz(name=name, wert=wert, beschriftung=beschriftung, einheit=einheit)
        )
        gelesen += 1

    for name, wert in ERGAENZTE_PARAMETER.items():
        beschriftung, einheit = PARAMETER_KATALOG[name]
        sitzung.merge(
            Parametersatz(name=name, wert=wert, beschriftung=beschriftung, einheit=einheit)
        )
        bericht.melde(f"Parameter '{name}' gibt es im Excel nicht, mit {wert} angelegt (F-3)")
        gelesen += 1
    bericht.zaehle("Parameter", gelesen)


def _maschinen(blatt, sitzung: Session, bericht: Importbericht, jahr: int) -> None:
    gelesen = 0
    for zeile in blatt.iter_rows(min_row=2, max_row=60, max_col=9, values_only=True):
        kuerzel = _text(zeile[0])
        if not kuerzel:
            continue
        wartung = int(_zahl(zeile[7]))
        sitzung.merge(
            Maschine(
                kuerzel=kuerzel,
                bezeichnung=_text(zeile[1]),
                verfuegbarkeit_prozent=_prozent(zeile[5]),
                wartung_jahr=jahr if wartung else None,
                wartung_woche=wartung or None,
                bemerkung=_text(zeile[8]),
            )
        )
        gelesen += 1
    bericht.zaehle("Maschinen", gelesen)


def _lieferanten(blatt, sitzung: Session, bericht: Importbericht) -> None:
    gelesen = 0
    for zeile in blatt.iter_rows(min_row=2, max_row=60, max_col=8, values_only=True):
        nummer = _text(zeile[0])
        if not nummer:
            continue
        bewertung = _text(zeile[4]).upper()
        if bewertung not in {"A", "B", "C"}:
            bericht.melde(f"Lieferant {nummer}: Bewertung '{bewertung}' ist unbekannt")
        sitzung.merge(
            Lieferant(
                liefnummer=nummer,
                name=_text(zeile[1]),
                ort=_text(zeile[2]),
                lieferzeit_wochen=int(_zahl(zeile[3])),
                bewertung=bewertung,
                erp_nummer=_text(zeile[5]),
                bemerkung=_text(zeile[7]),
            )
        )
        gelesen += 1
    bericht.zaehle("Lieferanten", gelesen)


def _rohmaterial(blatt, sitzung: Session, bericht: Importbericht) -> None:
    gelesen = 0
    for nummer_der_zeile, zeile in enumerate(
        blatt.iter_rows(min_row=2, max_row=200, max_col=10, values_only=True), start=2
    ):
        nummer = _text(zeile[0])
        if not nummer:
            continue
        if nummer_der_zeile > 16:
            bericht.melde(
                f"Rohmaterial {nummer} steht in Zeile {nummer_der_zeile}; das Blatt "
                "Bedarf kennt nur die Zeilen 2 bis 16, im Excel fehlte es still im "
                "Bestellvorschlag"
            )
        if not _text(zeile[8]):
            bericht.melde(f"Rohmaterial {nummer}: keine ERP-Nummer, Export überspringt es")
        sitzung.merge(
            Rohmaterial(
                rohnummer=nummer,
                bezeichnung=_text(zeile[1]),
                einheit=_text(zeile[2]),
                liefnummer=_text(zeile[3]),
                lager=_zahl(zeile[4]),
                bestellt_offen=_zahl(zeile[5]),
                gebinde=_zahl(zeile[6]),
                preis=_zahl(zeile[7]),
                erp_nummer=_text(zeile[8]),
                lagerort=_text(zeile[9]),
            )
        )
        gelesen += 1
    bericht.zaehle("Rohmaterial", gelesen)


def _artikel(blatt, sitzung: Session, bericht: Importbericht) -> set[str]:
    nummern: set[str] = set()
    for zeile in blatt.iter_rows(min_row=2, max_row=200, max_col=8, values_only=True):
        nummer = _text(zeile[0])
        if not nummer:
            continue
        material = _text(zeile[2])
        if not material_ist_bekannt(material):
            bericht.melde(
                f"Artikel {nummer}: Material '{material}' steht nicht in der Liste "
                "(R-010), unverändert übernommen"
            )
        sitzung.merge(
            Artikel(
                artikelnummer=nummer,
                bezeichnung=_text(zeile[1]),
                material=material,
                standardmaschine=_text(zeile[3]),
                minuten_je_stueck=_zahl(zeile[4]),
                ruestzeit_basis_min=int(_zahl(zeile[5])),
                zeichnung=_text(zeile[6]),
                aktiv=_text(zeile[7]).upper() != "N",
            )
        )
        nummern.add(nummer)
    bericht.zaehle("Artikel", len(nummern))
    return nummern


def _stueckliste(blatt, sitzung: Session, bericht: Importbericht, artikelnummern: set[str]) -> None:
    gelesen = 0
    je_artikel: dict[str, int] = {}
    for zeile in blatt.iter_rows(min_row=2, max_row=200, max_col=7, values_only=True):
        artikelnummer = _text(zeile[1])
        if not artikelnummer:
            continue
        if artikelnummer not in artikelnummern:
            bericht.melde(f"Stückliste: Artikel {artikelnummer} gibt es im Blatt Artikel nicht")
            continue
        position = int(_zahl(zeile[2]))
        je_artikel[artikelnummer] = je_artikel.get(artikelnummer, 0) + 1
        verschnitt = _prozent(zeile[5]) if _zahl(zeile[5]) <= 1 else _zahl(zeile[5])

        # Der Schlüssel der Zeile ist (Artikel, Position), nicht die laufende id.
        # Ohne diese Suche legt ein zweiter Import dieselbe Position noch einmal an.
        vorhanden = sitzung.scalar(
            select(Stuecklistenposition).where(
                Stuecklistenposition.artikelnummer == artikelnummer,
                Stuecklistenposition.position == position,
            )
        )
        if vorhanden is None:
            vorhanden = Stuecklistenposition(artikelnummer=artikelnummer, position=position)
            sitzung.add(vorhanden)
        vorhanden.rohnummer = _text(zeile[3])
        vorhanden.menge_je_stueck = _zahl(zeile[4])
        vorhanden.verschnitt_prozent = verschnitt
        vorhanden.bemerkung = _text(zeile[6])
        gelesen += 1
    for artikelnummer, anzahl in je_artikel.items():
        if anzahl > 2:
            bericht.melde(
                f"Artikel {artikelnummer} hat {anzahl} Positionen; das Excel konnte nur "
                "zwei rechnen (R-012, F-8), die App rechnet alle"
            )
    bericht.zaehle("Stücklistenpositionen", gelesen)


def _auftraege(
    blatt,
    sitzung: Session,
    bericht: Importbericht,
    artikelnummern: set[str],
    jahr: int,
) -> None:
    gelesen = 0
    kunden: dict[str, str] = {}
    for zeile in blatt.iter_rows(min_row=4, max_row=1000, max_col=18, values_only=True):
        nummer = _text(zeile[0])
        if not nummer:
            continue
        artikelnummer = _text(zeile[3])
        if artikelnummer not in artikelnummern:
            bericht.melde(
                f"Auftrag {nummer}: Artikel '{artikelnummer}' gibt es nicht, Zeile "
                "übersprungen (R-011)"
            )
            continue
        liefertermin = _datum(zeile[5])
        if liefertermin is None:
            bericht.melde(f"Auftrag {nummer}: kein Liefertermin, Zeile übersprungen")
            continue
        status = _text(zeile[16])
        if status not in {wert.value for wert in Status}:
            bericht.melde(
                f"Auftrag {nummer}: Status '{status}' ist keiner der drei erlaubten "
                "(R-008), Zeile übersprungen"
            )
            continue

        kunde = _text(zeile[1])
        if kunde and kunde not in kunden:
            kunden[kunde] = f"Kunde {len(kunden) + 1:02d}"

        manuell = int(_zahl(zeile[14]))
        auftrag = Auftrag(
            auftragsnummer=nummer,
            kunde=kunden.get(kunde, ""),
            prioritaet=_text(zeile[2]).upper(),
            artikelnummer=artikelnummer,
            menge=int(_zahl(zeile[4])),
            liefertermin=liefertermin,
            abweichende_maschine=None,
            status=status,
            bemerkung=_text(zeile[17]),
        )
        auftrag.setze_manuelle_kw(Kalenderwoche(jahr, manuell) if manuell else None)
        sitzung.merge(auftrag)
        gelesen += 1
    bericht.zaehle("Aufträge", gelesen)
    bericht.zaehle("Kunden anonymisiert", len(kunden))


def hauptprogramm(argumente: list[str] | None = None) -> int:
    zerleger = argparse.ArgumentParser(description="Excel in die Datenbank lesen")
    zerleger.add_argument("datei", type=Path, help="Pfad zur Arbeitsmappe (.xlsm)")
    zerleger.add_argument(
        "--trocken",
        action="store_true",
        help="nur lesen und prüfen, nichts in die Datenbank schreiben",
    )
    werte = zerleger.parse_args(argumente)

    if not werte.datei.exists():
        print(f"Datei nicht gefunden: {werte.datei}", file=sys.stderr)
        return 1

    tabellen_anlegen()
    with Sitzung() as sitzung:
        bericht = lies_excel(werte.datei, sitzung)
        if werte.trocken:
            sitzung.rollback()
            print("Trockenlauf, nichts geschrieben.")
        else:
            sitzung.commit()
    print(bericht.als_text())
    return 0


if __name__ == "__main__":
    raise SystemExit(hauptprogramm())
