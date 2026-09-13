"""Der Excel-Import liest die Eingabespalten und meldet Abweichungen, ohne zu korrigieren."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

from app.import_excel import lies_excel
from app.modelle import Artikel, Auftrag, Basis, Parametersatz, Rohmaterial


def _mappe(pfad: Path) -> Path:
    """Baut eine kleine Arbeitsmappe im Aufbau von Produktionsplanung_V7_3.xlsm."""
    mappe = Workbook()

    parameter = mappe.active
    parameter.title = "Parameter"
    parameter.append(["Bezeichnung", "Wert", "Bemerkung"])
    for bezeichnung, wert in [
        ("Aktuelle KW", 39),
        ("Jahr", 2026),
        ("Montag aktuelle KW", date(2026, 9, 21)),
        ("Sicherheitsbestand %", 0.10),
        ("Auslastungsgrenze", 0.90),
        ("Rüstzuschlag Alu (min)", 40),
        ("Losgrenze Alu (Stk)", 500),
        ("Schichten pro Tag", 2),
        ("Std pro Schicht", 8.0),
        ("Puffer Lieferant C (Wochen)", 1),
        ("Arbeitstage pro Woche", 5),
        ("Export-Pfad", "P:\\irgendwo"),
        ("Kostenstelle Einkauf", 4100),
        ("Planungshorizont (Wochen)", 8),
        ("Wartungsabzug (h)", 16),
        ("Zuschlag Edelstahl (min)", 20),
        ("Vorlauf Prio A (Wochen)", 2),
    ]:
        parameter.append([bezeichnung, wert])

    kapazitaet = mappe.create_sheet("Kapazität")
    kapazitaet.append(
        ["Maschine", "Bezeichnung", "Schichten", "Std", "Tage", "Verf.", "Kap.", "Wartung", "Bem."]
    )
    kapazitaet.append(["M1", "DMG CTX beta 800 (Drehen)", 2, 8.0, 5, 0.85, 68.0, 41, ""])
    kapazitaet.append(["M3", "Hermle C 22 (5-Achs)", 2, 8, 5, 0.80, 64.0, 0, "Einfahren"])

    lieferanten = mappe.create_sheet("Lieferanten")
    lieferanten.append(["LiefNr", "Name", "Ort", "Lieferzeit", "Bew.", "ERP", "Kontakt", "Bem."])
    lieferanten.append(["L-01", "Metallhandel Nord AG", "Zürich", 2, "A", "L01001", "", ""])
    lieferanten.append(["L-04", "Giesserei Ost AG", "Chur", 5, "C", "L04001", "", "immer +1 Wo"])

    roh = mappe.create_sheet("Rohmaterial")
    roh.append(
        [
            "RohNr",
            "Bezeichnung",
            "Einheit",
            "LiefNr",
            "Lager",
            "Bestellt",
            "Gebinde",
            "Preis",
            "ERP",
            "Ort",
        ]
    )
    roh.append(["RM-106", "Rundstange Alu Ø80", "Stange", "L-01", 5, 0, 5, 96.0, "3050800", "C-02"])
    roh.append(["RM-109", "Guss-Rohling Gehäuse", "Stk", "L-04", 90, 0, 20, 48.0, "", "D-01"])

    artikel = mappe.create_sheet("Artikel")
    artikel.append(
        ["ArtNr", "Bezeichnung", "Material", "Maschine", "min/Stk", "Rüst", "Zeichnung", "Aktiv"]
    )
    artikel.append(["A-1005", "Flansch DN50 Alu", "Alu", "M3", 4.0, 50, "Z-1005", "J"])
    artikel.append(["A-1009", "Gehäuse G12 GG25", "GG25", "M1", 18.0, 120, "Z-1009", "J"])
    artikel.append(["A-0998", "Welle alt", "ALU", "M1", 5.0, 30, "Z-0998", "N"])

    stueckliste = mappe.create_sheet("Stückliste")
    stueckliste.append(["Key", "ArtNr", "Pos", "RohNr", "Menge", "Verschnitt", "Bem."])
    stueckliste.append(["A-1005-1", "A-1005", 1, "RM-106", 0.0140, 0.08, ""])
    stueckliste.append(["A-1009-1", "A-1009", 1, "RM-109", 1.0, 0, ""])
    stueckliste.append(["A-0998-1", "A-0998", 1, "RM-106", 0.035, 0.05, "alt"])

    auftraege = mappe.create_sheet("Aufträge")
    auftraege.append(["Plan-KW: A = 2 Wochen vor Termin"])
    auftraege.append([])
    auftraege.append(
        [
            "AufNr",
            "Kunde",
            "Prio",
            "ArtNr",
            "Menge",
            "Termin",
            "Bez.",
            "Mat.",
            "Masch.",
            "Bearb",
            "Rüst",
            "Gesamt",
            "Liefer-KW",
            "Plan-KW",
            "Manuell",
            "KW eff.",
            "Status",
            "Bem.",
        ]
    )
    auftraege.append(
        [
            "26-0414",
            "Muster Werke AG",
            "B",
            "A-1005",
            620,
            date(2026, 10, 9),
            "",
            "",
            "",
            41.3,
            90,
            42.8,
            41,
            40,
            None,
            40,
            "offen",
            "",
        ]
    )
    auftraege.append(
        [
            "26-0420",
            "Beispiel GmbH",
            "C",
            "A-1009",
            30,
            date(2026, 10, 16),
            "",
            "",
            "",
            9.0,
            120,
            11.0,
            42,
            41,
            42,
            42,
            "offen",
            "auto verschoben",
        ]
    )
    auftraege.append(
        [
            "26-0421",
            "Muster Werke AG",
            "B",
            "A-9999",
            10,
            date(2026, 10, 16),
            "",
            "",
            "",
            0,
            0,
            0,
            42,
            41,
            None,
            41,
            "offen",
            "",
        ]
    )
    auftraege.append(
        [
            "26-0422",
            "Beispiel GmbH",
            "B",
            "A-1005",
            10,
            date(2026, 10, 16),
            "",
            "",
            "",
            0,
            0,
            0,
            42,
            41,
            None,
            41,
            "Offen",
            "",
        ]
    )

    ziel = pfad / "Produktionsplanung_demo.xlsx"
    mappe.save(ziel)
    return ziel


@pytest.fixture
def bericht_und_sitzung(tmp_path: Path):
    datei = _mappe(tmp_path)
    maschine = create_engine("sqlite://")
    Basis.metadata.create_all(maschine)
    fabrik = sessionmaker(bind=maschine, expire_on_commit=False)
    with fabrik() as sitzung:
        bericht = lies_excel(datei, sitzung)
        sitzung.commit()
        yield bericht, sitzung


def test_import_liest_stammdaten_und_auftraege(bericht_und_sitzung):
    bericht, sitzung = bericht_und_sitzung
    assert bericht.gelesen["Artikel"] == 3
    assert bericht.gelesen["Rohmaterial"] == 2
    assert bericht.gelesen["Aufträge"] == 2
    assert sitzung.scalar(select(Auftrag).where(Auftrag.auftragsnummer == "26-0414"))


def test_import_uebernimmt_nur_die_eingabespalten(bericht_und_sitzung):
    _, sitzung = bericht_und_sitzung
    auftrag = sitzung.get(Auftrag, "26-0414")
    assert auftrag.menge == 620
    assert auftrag.liefertermin == date(2026, 10, 9)
    # Die gerechneten Spalten des Excel stehen nicht in der Datenbank.
    assert not hasattr(auftrag, "gesamtzeit_h")


def test_import_liest_die_manuelle_woche_mit_jahr(bericht_und_sitzung):
    _, sitzung = bericht_und_sitzung
    auftrag = sitzung.get(Auftrag, "26-0420")
    assert auftrag.manuelle_kw.jahr == 2026
    assert auftrag.manuelle_kw.woche == 42


def test_import_anonymisiert_die_kunden(bericht_und_sitzung):
    _, sitzung = bericht_und_sitzung
    for auftrag in sitzung.scalars(select(Auftrag)).all():
        assert auftrag.kunde.startswith("Kunde ")


def test_import_meldet_falsche_materialschreibweise(bericht_und_sitzung):
    bericht, sitzung = bericht_und_sitzung
    assert any("ALU" in meldung for meldung in bericht.meldungen)
    # Er korrigiert nicht still.
    assert sitzung.get(Artikel, "A-0998").material == "ALU"


def test_import_ueberspringt_unbekannte_artikel_und_status(bericht_und_sitzung):
    bericht, sitzung = bericht_und_sitzung
    assert sitzung.get(Auftrag, "26-0421") is None
    assert sitzung.get(Auftrag, "26-0422") is None
    assert any("A-9999" in meldung for meldung in bericht.meldungen)
    assert any("'Offen'" in meldung for meldung in bericht.meldungen)


def test_import_meldet_fehlende_erp_nummer(bericht_und_sitzung):
    bericht, _ = bericht_und_sitzung
    assert any("RM-109" in meldung and "ERP" in meldung for meldung in bericht.meldungen)


def test_import_rechnet_prozentzellen_um(bericht_und_sitzung):
    _, sitzung = bericht_und_sitzung
    werte = {satz.name: satz.wert for satz in sitzung.scalars(select(Parametersatz)).all()}
    assert werte["auslastungsgrenze_prozent"] == Decimal(90)
    assert werte["sicherheitsbestand_prozent"] == Decimal(10)


def test_import_ergaenzt_die_fehlenden_vorlauf_parameter(bericht_und_sitzung):
    bericht, sitzung = bericht_und_sitzung
    werte = {satz.name: satz.wert for satz in sitzung.scalars(select(Parametersatz)).all()}
    assert werte["vorlauf_prio_b_wochen"] == Decimal(1)
    assert werte["vorlauf_prio_c_wochen"] == Decimal(1)
    assert any("F-3" in meldung for meldung in bericht.meldungen)


def test_import_uebernimmt_export_pfad_und_aktuelle_kw_nicht(bericht_und_sitzung):
    _, sitzung = bericht_und_sitzung
    namen = {satz.name for satz in sitzung.scalars(select(Parametersatz)).all()}
    assert "export_pfad" not in namen
    assert "aktuelle_kw" not in namen


def test_import_liest_verschnitt_und_stueckliste(bericht_und_sitzung):
    _, sitzung = bericht_und_sitzung
    artikel = sitzung.get(Artikel, "A-1005")
    position = artikel.positionen[0]
    assert position.rohnummer == "RM-106"
    assert position.menge_je_stueck == Decimal("0.014")
    assert position.verschnitt_prozent == Decimal(8)
    assert sitzung.get(Rohmaterial, "RM-106").gebinde == Decimal(5)
